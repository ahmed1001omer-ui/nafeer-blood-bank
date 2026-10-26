from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

ROOT=Path.cwd()
OUT=ROOT/'نموذج تجريبي - إدارة بنك الدم.docx'
AS=ROOT/'draft_assets'
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(2.5); sec.bottom_margin=Cm(2.3); sec.left_margin=Cm(2.7); sec.right_margin=Cm(2.7)
sec.header_distance=Cm(1.2); sec.footer_distance=Cm(1.2)

# Set Arabic-friendly font and paragraph direction.
def set_run_font(run, size=14, bold=False, name='Simplified Arabic'):
    run.font.name=name; run.font.size=Pt(size); run.bold=bold
    rpr=run._element.get_or_add_rPr()
    rf=rpr.rFonts
    if rf is not None:
        rf.set(qn('w:ascii'),name); rf.set(qn('w:hAnsi'),name); rf.set(qn('w:eastAsia'),name); rf.set(qn('w:cs'),name)
    return run

def rtl(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.alignment=align
    ppr=p._p.get_or_add_pPr(); bidi=OxmlElement('w:bidi'); bidi.set(qn('w:val'),'1'); ppr.append(bidi)
    p.paragraph_format.line_spacing=1.5
    p.paragraph_format.space_after=Pt(7)
    return p

def addp(text='', size=14, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=7, first_indent=None):
    p=doc.add_paragraph(); rtl(p,align)
    p.paragraph_format.space_before=Pt(before); p.paragraph_format.space_after=Pt(after)
    if first_indent is not None: p.paragraph_format.first_line_indent=Cm(first_indent)
    set_run_font(p.add_run(text),size,bold)
    return p

def heading(text, level=1):
    size={1:16,2:14,3:14}.get(level,14)
    p=doc.add_paragraph(); rtl(p,WD_ALIGN_PARAGRAPH.RIGHT)
    p.paragraph_format.keep_with_next=True; p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(5)
    set_run_font(p.add_run(text),size,True)
    # Native heading style for navigation and automatic TOC compatibility.
    p.style=f'Heading {level}'
    for r in p.runs: set_run_font(r,size,True)
    rtl(p,WD_ALIGN_PARAGRAPH.RIGHT)
    return p

def cell_fill(cell, color='D9E2F3'):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)

def set_cell_borders(table):
    tblPr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'5'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    tblPr.append(borders)

def field(run, instr):
    begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    ins=OxmlElement('w:instrText'); ins.set(qn('xml:space'),'preserve'); ins.text=instr
    sep=OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'),'separate')
    txt=OxmlElement('w:t'); txt.text='1'
    end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.append(begin); run._r.append(ins); run._r.append(sep); run._r.append(txt); run._r.append(end)

# Default body style.
normal=doc.styles['Normal']; normal.font.name='Simplified Arabic'; normal.font.size=Pt(14)
normal._element.rPr.rFonts.set(qn('w:cs'),'Simplified Arabic')
for sname in ('Heading 1','Heading 2','Heading 3'):
    st=doc.styles[sname]; st.font.name='Simplified Arabic'; st.font.color.rgb=None
    st._element.rPr.rFonts.set(qn('w:cs'),'Simplified Arabic')

# Page 1 cover with university marks from the supplied university guide.
t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for c,w in zip(t.rows[0].cells,[Cm(3.0),Cm(10.3),Cm(3.0)]): c.width=w; c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for idx,f in [(0,AS/'faculty-seal.jpg'),(2,AS/'university-seal.jpg')]:
    p=t.cell(0,idx).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    if f.exists(): p.add_run().add_picture(str(f),width=Cm(2.25))
p=t.cell(0,1).paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
for i,txt in enumerate(['جامعة النيل الأبيض','كلية علوم الحاسوب وتقنية المعلومات','قسم تقنية المعلومات / نظم المعلومات']):
    if i: p.add_run('\n')
    set_run_font(p.add_run(txt),15 if i<2 else 14,True)
set_cell_borders(t)
addp('بسم الله الرحمن الرحيم',16,True,WD_ALIGN_PARAGRAPH.CENTER,before=20,after=18)
addp('بحث تكميلي لنيل درجة البكالوريوس في تقنية المعلومات / نظم المعلومات',14,True,WD_ALIGN_PARAGRAPH.CENTER,after=25)
addp('تصميم وتنفيذ نظام إلكتروني لإدارة بنك الدم',16,True,WD_ALIGN_PARAGRAPH.CENTER,after=6)
addp('نموذج تجريبي وفق توجيهات الجامعة',14,True,WD_ALIGN_PARAGRAPH.CENTER,after=40)
addp('إعداد الطالب / الطالبة:',16,True,WD_ALIGN_PARAGRAPH.RIGHT,after=5)
addp('اسم افتراضي للطالب الأول (يستبدل بالاسم الحقيقي)',14,False,WD_ALIGN_PARAGRAPH.RIGHT,after=3)
addp('اسم افتراضي للطالب الثاني (يستبدل بالاسم الحقيقي)',14,False,WD_ALIGN_PARAGRAPH.RIGHT,after=25)
addp('إشراف الأستاذ / ................................................',16,True,WD_ALIGN_PARAGRAPH.RIGHT,after=35)
addp('العام الجامعي: ................................',14,False,WD_ALIGN_PARAGRAPH.CENTER,after=0)
doc.add_page_break()

# Preliminary matter.
heading('الآية القرآنية',1)
addp('﴿وَمَنْ أَحْيَاهَا فَكَأَنَّمَا أَحْيَا النَّاسَ جَمِيعًا﴾',18,True,WD_ALIGN_PARAGRAPH.CENTER,before=55,after=12)
addp('سورة المائدة، الآية 32',14,False,WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()
heading('الإهداء',1)
addp('نهدي هذا العمل إلى أسرنا، وإلى كل من شجعنا وساندنا في مسيرتنا العلمية، وإلى العاملين في خدمات نقل الدم الذين يسهمون في تنظيم العمل وخدمة المرضى.',16,False,WD_ALIGN_PARAGRAPH.CENTER,before=70,after=20)
doc.add_page_break()
heading('الشكر والتقدير',1)
addp('الحمد لله الذي وفقنا لإعداد هذا النموذج. ونتقدم بالشكر إلى جامعة النيل الأبيض وكلية علوم الحاسوب وتقنية المعلومات، وإلى مشرف البحث على ما سيقدمه من توجيه وملاحظات عند استكمال المشروع. كما نقدر جهود العاملين في خدمات بنك الدم، ونؤكد أن هذا النص نموذج أولي وأن أسماء الأشخاص والجهات ستثبت بعد اعتمادها.',14,False,WD_ALIGN_PARAGRAPH.JUSTIFY,before=10,after=14)
doc.add_page_break()
heading('الملخص',1)
addp('تتناول هذه الدراسة المقترحة تصميم وتنفيذ نموذج أولي لنظام إلكتروني يساعد على تنظيم إجراءات بنك الدم، مع التركيز على تسجيل المتبرعين، وإحالة بياناتهم إلى المختبر، وإدارة نتائج الفحوص، ومتابعة وحدات الدم ومخزونها. تنطلق المشكلة من الحاجة إلى ربط المعلومات والإجراءات بين الاستقبال والمختبر والمخزون، بما يقلل تكرار إدخال البيانات ويسهّل الوصول إلى السجلات ومتابعة حالة الوحدة. يهدف المشروع إلى تحليل متطلبات النظام، وبناء نموذج تجريبي يضم صلاحيات وأدوارًا وظيفية، وسجلًا للمتبرعين، وسير عمل لإحالة الفحوص، وشاشة للمخزون والتنبيهات. يعتمد العمل منهجًا وصفيًا تحليليًا في تحديد المتطلبات، ومنهجًا تطبيقيًا في تصميم النموذج وتنفيذه. لا يفترض هذا النموذج اتصالًا فعليًا بخادم مركزي أو تطبيق التبرع القائم، ولا يتخذ قرارًا طبيًا مستقلًا بشأن أهلية المتبرع؛ إذ يجب أن تستند هذه القرارات إلى بروتوكولات الجهة الصحية المختصة ومراجعة المختصين. تمثل الوثيقة نموذجًا أوليًا يحتاج إلى اعتماد عنوانه ونطاقه، والتحقق الميداني من المتطلبات، ثم تنفيذ النظام وتقييمه قبل استخلاص النتائج النهائية.',14,False)
addp('الكلمات المفتاحية: بنك الدم، إدارة المتبرعين، نظام معلومات، المختبر، إدارة المخزون، التتبع.',12,True,WD_ALIGN_PARAGRAPH.RIGHT,after=0)
doc.add_page_break()
heading('Abstract',1)
addp('This proposed study concerns the design and implementation of a prototype information system for organizing blood bank workflows. It focuses on donor registration, referral of donor records to the laboratory, management of screening results, and tracking blood units and inventory. The problem addressed is the need to connect reception, laboratory, and inventory information so that staff can retrieve records and follow the status of donations with less duplicated entry. The project aims to analyze requirements and build a prototype with role-based access, donor records, a laboratory referral workflow, and inventory views and alerts. A descriptive and analytical approach will be used to identify requirements, followed by an applied system development approach. The prototype is not assumed to connect to a production server or to the existing donor application, and it will not make independent medical eligibility decisions. Such decisions must follow the competent health authority’s protocols and qualified clinical review. This document is an initial academic sample; its title and scope, field requirements, implementation, and evaluation remain subject to confirmation and completion.',14,False,WD_ALIGN_PARAGRAPH.JUSTIFY)
addp('Keywords: blood bank, donor management, information system, laboratory, inventory management, traceability.',12,True,WD_ALIGN_PARAGRAPH.LEFT,after=0)
doc.add_page_break()

heading('قائمة المحتويات',1)
contents=[
('الآية القرآنية','أ'),('الإهداء','ب'),('الشكر والتقدير','ج'),('الملخص','د'),('Abstract','هـ'),
('الباب الأول: الإطار العام للبحث','1'),('1.1 المقدمة','1'),('1.2 مشكلة البحث','2'),('1.3 أسئلة البحث','2'),('1.4 فرضيات البحث','2'),('1.5 أهداف البحث','3'),('1.6 أهمية البحث','3'),('1.7 حدود البحث','4'),('1.8 منهجية البحث','4'),('1.9 مصطلحات ومفاهيم البحث','5'),('1.10 هيكلية البحث','6'),('المراجع','7')]
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
t.columns[0].width=Cm(13.5); t.columns[1].width=Cm(2.5)
for c,txt in zip(t.rows[0].cells,['الموضوع','الصفحة']):
    c.text=txt; cell_fill(c,'1F4E78'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in c.paragraphs:
        rtl(p,WD_ALIGN_PARAGRAPH.CENTER)
        for r in p.runs: set_run_font(r,12,True); r.font.color.rgb=__import__('docx').shared.RGBColor(255,255,255)
for label,page in contents:
    cells=t.add_row().cells; cells[0].text=label; cells[1].text=page
    for i,c in enumerate(cells):
        c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in c.paragraphs:
            rtl(p,WD_ALIGN_PARAGRAPH.RIGHT if i==0 else WD_ALIGN_PARAGRAPH.CENTER); p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.15
            for r in p.runs: set_run_font(r,11,False)
set_cell_borders(t)
addp('ملاحظة: أرقام الصفحات أعلاه تجريبية، وتُحدّث بعد اعتماد المحتوى النهائي وإعادة إنشاء الفهرس.',10,False,WD_ALIGN_PARAGRAPH.RIGHT,before=6,after=0)
doc.add_page_break()

# Chapter 1 begins a new numbering section.
chapter_sec=doc.add_section(WD_SECTION.NEW_PAGE)
chapter_sec.page_width=Cm(21); chapter_sec.page_height=Cm(29.7)
chapter_sec.top_margin=Cm(2.5); chapter_sec.bottom_margin=Cm(2.3); chapter_sec.left_margin=Cm(2.7); chapter_sec.right_margin=Cm(2.7)
chapter_sec.header_distance=Cm(1.2); chapter_sec.footer_distance=Cm(1.2)
chapter_sec.footer.is_linked_to_previous=False
pgn=OxmlElement('w:pgNumType'); pgn.set(qn('w:start'),'1'); chapter_sec._sectPr.append(pgn)
heading('الباب الأول: الإطار العام للبحث',1)
heading('1.1 المقدمة',2)
addp('تتطلب خدمات نقل الدم تنسيقًا دقيقًا للمعلومات والإجراءات منذ تسجيل المتبرع وجمع بيانات التبرع، مرورًا بالفحوص المخبرية، ووصولًا إلى إدارة مكونات الدم وحفظ سجلاتها. وتزداد أهمية هذا التنسيق عندما تتوزع المهام على وحدات مختلفة، مثل الاستقبال والمختبر والمخزون، لأن انقطاع انتقال البيانات أو تأخر تحديثها قد يصعّب متابعة السجل والوحدة المرتبطة به.',14,False,first_indent=0.7)
addp('تؤكد إرشادات منظمة الصحة العالمية أهمية وجود مسار معلومات موثوق يربط المتبرع بالتبرع ومكوناته، مع الاحتفاظ بالسجلات اللازمة للتتبع. كما أن إدارة سلسلة التبريد وحالة التخزين وتواريخ الجمع والانتهاء عناصر مؤثرة في صلاحية المكونات وسلامتها [1]، [2]. ومن ثم يمكن لنظام معلومات مناسب أن يدعم حفظ البيانات، وتبادلها بين الموظفين المخولين، وإظهار معلومات المخزون وحالته، مع المحافظة على أن القرارات المخبرية والسريرية تقع ضمن صلاحيات المختصين.',14,False,first_indent=0.7)
addp('يقترح هذا البحث نموذجًا إلكترونيًا لإدارة سير العمل في بنك الدم، يضم تسجيل المتبرع في الاستقبال، وإرسال السجل إلى المختبر لاستكمال البيانات والنتائج، ومتابعة وحدات الدم والمخزون، وإدارة حسابات الموظفين وصلاحياتهم. ويقتصر النموذج الأولي على بيئة تجريبية وبيانات اختبار غير حقيقية، ولا يشمل بناء تطبيق التبرع القائم أو الربط بخادم إنتاجي ضمن النطاق الحالي.',14,False,first_indent=0.7)
heading('1.2 مشكلة البحث',2)
addp('تتمثل مشكلة البحث في احتمال تشتت بيانات المتبرعين ونتائج الفحوص ومعلومات وحدات الدم بين إجراءات أو سجلات منفصلة، وما قد ينتج عن ذلك من تكرار إدخال البيانات، وصعوبة معرفة حالة السجل، وتأخر انتقال المعلومات من الاستقبال إلى المختبر، أو محدودية متابعة المخزون. وتبين وثيقة المتطلبات الأولية المرفقة حاجة المستخدم إلى سجل متبرعين يسجل فيه الموظف البيانات، وإحالة هذه البيانات إلى المختبر، وإلى عرض مخزون الدم المتاح مع معلومات الحالة والانتهاء والموقع.',14,False,first_indent=0.7)
addp('وعليه تتمحور المشكلة في السؤال الآتي: كيف يمكن تصميم نموذج نظام معلومات يربط إجراءات الاستقبال والمختبر والمخزون في بنك الدم، وينظم تسجيل البيانات وتتبعها وعرضها، مع تحديد الصلاحيات والمحافظة على أن قرار القبول الطبي من مسؤولية المختص؟',14,False,first_indent=0.7)
heading('1.3 أسئلة البحث',2)
for x in [
'ما البيانات والوظائف اللازمة لتنظيم تسجيل المتبرع وإحالة سجله من الاستقبال إلى المختبر؟',
'كيف يمكن تنظيم عرض وحدات الدم ومخزونها وحالاتها وتواريخها بما يسهل استرجاع المعلومات؟',
'ما الأدوار والصلاحيات المطلوبة لموظفي الاستقبال والمختبر والإدارة؟',
'كيف يمكن بناء نموذج أولي قابل للتقييم دون ربطه بخادم إنتاجي أو اتخاذه قرارًا طبيًا مستقلًا؟']:
    addp('• '+x,14,False,WD_ALIGN_PARAGRAPH.RIGHT,after=3)
heading('1.4 فرضيات البحث',2)
addp('لا يضع هذا النموذج فرضيات إحصائية؛ فالعمل المقترح تطبيقي يركز على تحليل المتطلبات وتصميم نظام أولي. ويمكن لاحقًا إضافة فرضيات قابلة للقياس إذا تضمنت الدراسة مقارنة كمية، مثل قياس زمن استرجاع السجل قبل النظام وبعده، شريطة جمع بيانات فعلية واعتماد خطة التقييم.',14,False,first_indent=0.7)
heading('1.5 أهداف البحث',2)
heading('الهدف العام',3)
addp('تصميم وتنفيذ نموذج أولي لنظام إلكتروني ينظم إجراءات الاستقبال والمختبر ومخزون الدم، ويدعم حفظ المعلومات وتتبعها ضمن صلاحيات وظيفية محددة.',14,False,first_indent=0.7)
heading('الأهداف الفرعية',3)
for x in [
'تحليل متطلبات تسجيل المتبرعين وإدارة ملفاتهم في الاستقبال.',
'تصميم آلية إحالة بيانات المتبرع إلى المختبر واستكمال النتائج في السجل نفسه.',
'تصميم شاشة للمخزون تعرض الفصيلة والمكون والحالة وتاريخي السحب والانتهاء وموقع التخزين وفق البيانات المتاحة.',
'تحديد الأدوار والصلاحيات الأساسية للموظفين والإدارة.',
'إنشاء نموذج تجريبي ببيانات افتراضية، ثم إعداد خطة لتقييم وظائفه مع المستخدمين المختصين.',
'توثيق حدود النموذج وما يتطلبه الانتقال مستقبلًا إلى تشغيل فعلي، مثل اعتماد البروتوكولات الصحية وأمن البيانات والتكامل مع الأنظمة القائمة.']:
    addp('• '+x,14,False,WD_ALIGN_PARAGRAPH.RIGHT,after=3)
heading('1.6 أهمية البحث',2)
heading('الأهمية العلمية',3)
addp('يسهم البحث في تطبيق مفاهيم تحليل وتصميم نظم المعلومات على بيئة عمل متعددة الأدوار، ويوضح كيفية تحويل احتياجات الاستقبال والمختبر والمخزون إلى وظائف وسجلات مترابطة قابلة للتقييم. كما يتيح أساسًا لدراسات لاحقة حول جودة البيانات، وقابلية الاستخدام، والتتبع في أنظمة خدمات نقل الدم.',14,False,first_indent=0.7)
heading('الأهمية العملية والتطبيقية',3)
addp('قد يساعد النموذج المقترح على تنظيم إدخال البيانات، وإظهار حالة السجل بين الأقسام، وتسهيل الوصول إلى معلومات وحدات الدم المخزنة. وتظل الفائدة التشغيلية الفعلية مرتبطة بمراجعة المتطلبات مع بنك دم معتمد، والتحقق من الأمن والخصوصية، واعتماد سياسات التخزين والفحوص والأهلية من الجهات الصحية المختصة قبل أي استخدام حقيقي.',14,False,first_indent=0.7)
heading('1.7 حدود البحث',2)
heading('الحدود الموضوعية',3)
addp('يركز النموذج على إدارة سجلات المتبرعين، وإحالة البيانات إلى المختبر، وتسجيل النتائج من قبل الموظف المخول، وعرض وحدات الدم والمخزون، وإدارة المستخدمين والصلاحيات. ولا يتضمن إصدار حكم طبي آلي على أهلية التبرع أو استبدال إجراءات الفحص المعتمدة.',14,False)
heading('الحدود المكانية',3)
addp('يُحدد موقع التطبيق بعد التنسيق مع جهة بنك الدم التي ستشارك في تحليل المتطلبات. ولم يحدد الملف المرفق مؤسسة ميدانية بعينها، لذلك لا ينسب هذا النموذج إلى مركز أو مستشفى محدد.',14,False)
heading('الحدود الزمانية',3)
addp('تغطي هذه الوثيقة مرحلة إعداد النموذج الأكاديمي الأولي. وتضاف فترة التنفيذ وجمع البيانات والتقييم بعد اعتماد الخطة والجدول الزمني من القسم والمشرف.',14,False)
heading('الحدود البشرية',3)
addp('الفئات المستفيدة المقترحة هي موظفو الاستقبال والمختبر وإدارة بنك الدم. وتحدد العينة الفعلية وعدد المشاركين بعد موافقة الجهة المعنية واعتماد إجراءات جمع البيانات.',14,False)
heading('1.8 منهجية البحث',2)
heading('منهج البحث',3)
addp('يقترح البحث الجمع بين المنهج الوصفي التحليلي لفهم الإجراءات والمتطلبات، والمنهج التطبيقي لتصميم النموذج وتنفيذه. وتتضمن الخطة المبدئية: تحديد أصحاب المصلحة، توثيق المتطلبات الوظيفية وغير الوظيفية، إعداد نماذج العمليات والبيانات، بناء النموذج الأولي، ثم تقييمه وفق سيناريوهات استخدام متفق عليها. لم تنفذ هذه الخطوات الميدانية بعد في هذا النموذج.',14,False)
heading('مجتمع البحث',3)
addp('مجتمع البحث المقترح هو العاملون في أقسام الاستقبال والمختبر وإدارة المخزون في جهة بنك الدم التي ستعتمدها الدراسة.',14,False)
heading('عينة البحث',3)
addp('تحدد العينة لاحقًا بالتنسيق مع الجهة الميدانية. ويمكن أن تكون عينة قصدية من الموظفين الذين يباشرون الإجراءات التي يغطيها النظام، بعد الموافقات اللازمة.',14,False)
heading('أدوات جمع البيانات',3)
addp('تقترح الدراسة مقابلات أو استبانة قصيرة ومراجعة النماذج والإجراءات المعتمدة، إلى جانب وثيقة المتطلبات المرفقة. ولا تعد المتطلبات الواردة في الوثيقة وحدها دليلًا على اعتمادها من جهة صحية أو جامعية.',14,False)
heading('أساليب تحليل البيانات',3)
addp('تحلل الملاحظات والمقابلات موضوعيًا لاستخراج الأدوار والعمليات والبيانات المطلوبة، ثم تصنف المتطلبات إلى وظيفية وغير وظيفية. ويقيّم النموذج لاحقًا بسيناريوهات وظيفية وملاحظات المستخدمين؛ ولا تعرض هذه الوثيقة نتائج تقييم لم يحدث بعد.',14,False)
heading('1.9 مصطلحات ومفاهيم البحث',2)
for term,desc in [
('نظام معلومات بنك الدم','نظام إلكتروني مقترح لتنظيم بيانات المتبرعين وإجراءات الأقسام ومعلومات وحدات الدم ضمن صلاحيات محددة.'),
('المتبرع','الشخص الذي يسجل بياناته وتبرعه في النظام، على أن تحدد الجهة المختصة ضوابط التبرع وتقييم الأهلية.'),
('إحالة المختبر','انتقال سجل المتبرع من مرحلة الاستقبال إلى قائمة عمل المختبر لاستكمال الفحوص وتوثيق نتائجها بواسطة موظف مخول.'),
('وحدة الدم ومكوناتها','سجل تعريفي لوحدة أو مكون ناتج عن التبرع، يرتبط ببيانات التعريف والحالة والتواريخ ومعلومات التخزين التي تعتمدها الجهة.'),
('التتبع','إمكانية ربط المعلومات ذات الصلة بالمتبرع والتبرع والمكونات ومراحل التعامل معها، وفق سجلات وصلاحيات تحفظ دقة البيانات.'),
('النموذج الأولي','نسخة تجريبية من النظام تستخدم للتحقق من الوظائف والتصميم، ولا تعني جاهزيتها للتشغيل السريري أو التجاري.')]:
    p=doc.add_paragraph(); rtl(p,WD_ALIGN_PARAGRAPH.JUSTIFY); p.paragraph_format.space_after=Pt(6)
    set_run_font(p.add_run(term+': '),14,True); set_run_font(p.add_run(desc),14,False)
heading('1.10 هيكلية البحث',2)
addp('يقترح تنظيم البحث النهائي في أبواب مترابطة. يعرض الباب الأول الإطار العام ومشكلة البحث وأهدافه ومنهجيته وحدوده. ويتناول الباب الثاني الخلفية النظرية والدراسات السابقة ومفاهيم خدمات بنك الدم ونظم المعلومات ذات الصلة. ويعرض الباب الثالث تحليل النظام وتصميمه، بما في ذلك المتطلبات والعمليات والبيانات والواجهات. ويشرح الباب الرابع التنفيذ والتقييم والنتائج، ثم يتضمن البحث خاتمة وتوصيات ومصادر وملاحق وفق توجيهات الكلية. هذه الهيكلية مقترح أولي وتعدل حسب اعتماد القسم والمشرف.',14,False,first_indent=0.7)

heading('المراجع',1)
addp('[1] World Health Organization, “A guide to establishing a national haemovigilance system,” WHO, 2016. [Online]. Available: https://iris.who.int/bitstream/handle/10665/250233/9789241549844-eng.pdf. [Accessed: Oct. 7, 2026].',11,False,WD_ALIGN_PARAGRAPH.LEFT,after=6)
addp('[2] World Health Organization, “Blood cold chain,” Health Products Policy and Standards. [Online]. Available: https://www.who.int/teams/health-product-policy-and-standards/standards-and-specifications/blood-and-products-of-human-origin/blood-cold-chain. [Accessed: Oct. 7, 2026].',11,False,WD_ALIGN_PARAGRAPH.LEFT,after=6)
addp('ملاحظة توثيقية: أدرجت المراجع للاستناد إلى أهمية التتبع وسلسلة التبريد فقط؛ ويجب إضافة مراجع متخصصة أخرى بعد اعتماد نطاق البحث ومراجعته مع المشرف.',11,True,WD_ALIGN_PARAGRAPH.RIGHT,before=4,after=0)

# Front matter uses alphabetic page fields; chapter numbering restarts at 1.
doc.sections[0].different_first_page_header_footer=True
fp=doc.sections[0].footer.paragraphs[0]; rtl(fp,WD_ALIGN_PARAGRAPH.CENTER); fp.paragraph_format.space_after=Pt(0)
set_run_font(fp.add_run('— '),10,False); rr=fp.add_run(); set_run_font(rr,10,False); field(rr,'PAGE \\* alphabetic'); set_run_font(fp.add_run(' —'),10,False)
fp2=doc.sections[1].footer.paragraphs[0]; rtl(fp2,WD_ALIGN_PARAGRAPH.CENTER); fp2.paragraph_format.space_after=Pt(0)
set_run_font(fp2.add_run('— '),10,False); rr=fp2.add_run(); set_run_font(rr,10,False); field(rr,'PAGE \\* Arabic'); set_run_font(fp2.add_run(' —'),10,False)

# Update fields when opened in Word.
settings=doc.settings.element
uf=OxmlElement('w:updateFields'); uf.set(qn('w:val'),'true'); settings.append(uf)
doc.core_properties.title='تصميم وتنفيذ نظام إلكتروني لإدارة بنك الدم — نموذج تجريبي'
doc.core_properties.subject='نموذج بحث وفق توجيهات جامعة النيل الأبيض'
doc.core_properties.author='نموذج افتراضي للاختبار'
doc.save(OUT)
print(OUT)

